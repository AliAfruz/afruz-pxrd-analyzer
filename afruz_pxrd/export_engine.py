from __future__ import annotations

"""Phase 22 Raptor export engine.

The engine converts the complete scientific state into explicitly named,
rectangular tables before writing either a styled Excel workbook or a clean
UTF-8 TXT package.  It deliberately avoids dumping nested dictionaries into
unreadable cells whenever a result can be represented as a summary, profile,
phase, reflection, peak, or parameter table.
"""

from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
import hashlib
import json
import math
import os
import re
import shutil
import tempfile
import zipfile
from typing import Any

try:
    import numpy as np
except Exception:  # pragma: no cover
    np = None

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from .contracts.base import decode_value
from .text_export import atomic_write_text, safe_filename, sha256_file, write_manifest_txt, write_table_txt
from .manuscript_plot_export import MANUSCRIPT_COLUMN_SPECS, build_manuscript_plot_data
from .version import APP_VERSION

EXPORT_SCHEMA_VERSION = 2
MISSING_TEXT = "NA"

CATEGORY_LABELS = {
    "project": "Project and dataset",
    "pattern": "Diffraction pattern",
    "preprocessing": "Preprocessing",
    "peaks": "Peaks and fitting",
    "structure": "Phase and structure",
    "refinement": "Whole-pattern refinement",
    "qpa": "Quantitative phase analysis",
    "validation": "Validation and audit",
    "batch": "Batch analysis",
}

# The normal export surface is deliberately small.  The broad CATEGORY_LABELS
# map remains available through the advanced/audit export so no reproducibility
# information is lost.
CLEAN_CONTENT_LABELS = {
    "raw": "Raw data",
    "treated": "Treated data",
    "peaks": "Peaks",
    "refinement": "Refinements — Pawley, Le Bail and Rietveld",
}

COLUMN_LABELS: dict[str, tuple[str, str, str]] = {
    "point_index": ("Point index", "", "Sequential row number beginning at 1."),
    "two_theta_deg": ("2θ", "°", "Diffraction angle in degrees two-theta."),
    "two_theta": ("2θ", "°", "Diffraction angle in degrees two-theta."),
    "position": ("Peak position", "° 2θ", "Observed or fitted peak-center position."),
    "center": ("Peak center", "° 2θ", "Peak-center position."),
    "observed_intensity": ("Observed intensity", "counts or input units", "Measured intensity used by the analysis."),
    "intensity_raw": ("Raw intensity", "counts or input units", "Imported intensity before processing."),
    "intensity_processed": ("Processed intensity", "processed units", "Intensity after the selected preprocessing operations."),
    "intensity_treated": ("Treated intensity", "processed units", "Final intensity after the selected preprocessing operations."),
    "intensity": ("Intensity", "counts or input units", "Peak or reflection intensity."),
    "fwhm_deg": ("FWHM", "° 2θ", "Full width at half maximum."),
    "calculated_intensity": ("Calculated intensity", "observed units", "Model-calculated profile in the same intensity units as the observation."),
    "background_intensity": ("Background", "observed units", "Calculated or estimated background profile."),
    "difference_observed_minus_calculated": ("Difference: observed − calculated", "observed units", "Point-wise residual profile."),
    "observed_sigma": ("Observed σ", "observed units", "Standard uncertainty assigned to the observed intensity."),
    "statistical_weight": ("Statistical weight", "1 / intensity²", "Weight used for statistically valid residual statistics."),
    "h": ("h", "", "Miller index h."),
    "k": ("k", "", "Miller index k."),
    "l": ("l", "", "Miller index l."),
    "hkl_label": ("Reflection", "hkl", "Reflection label."),
    "d_spacing": ("d spacing", "Å", "Interplanar spacing."),
    "d_spacing_angstrom": ("d spacing", "Å", "Interplanar spacing."),
    "fwhm": ("FWHM", "° 2θ", "Full width at half maximum."),
    "area": ("Integrated area", "intensity·° 2θ", "Integrated peak area."),
    "height": ("Peak height", "intensity", "Peak amplitude above the fitted baseline."),
    "rwp_percent": ("Rwp", "%", "Weighted profile residual."),
    "rp_percent": ("Rp", "%", "Unweighted profile residual."),
    "rexp_percent": ("Rexp", "%", "Expected profile residual when statistical weights are valid."),
    "goodness_of_fit_sqrt": ("GoF", "", "Rwp divided by Rexp."),
    "reduced_chi_square": ("Reduced χ²", "", "Square of the goodness of fit."),
    "weight_fraction_percent": ("Weight fraction", "wt%", "Refined quantitative phase fraction when scientifically valid."),
    "phase_fraction_percent": ("Phase fraction", "%", "Reported phase contribution or fraction."),
    "phase_name": ("Phase", "", "Phase or material name."),
    "formula": ("Formula", "", "Chemical formula from the structure or result."),
    "source_path": ("Source file", "", "Original imported file path."),
    "dataset_uid": ("Dataset UID", "", "Stable internal dataset identifier."),
    "dataset_name": ("Dataset", "", "Dataset display name."),
    "included": ("Included", "yes/no", "Whether the row was included in the calculation."),
    "use": ("Included", "yes/no", "Whether the peak was included in downstream calculations."),
    "manual": ("Manual peak", "yes/no", "Whether the peak was entered manually."),
    "uncertainty": ("Standard uncertainty", "", "Estimated standard uncertainty."),
    "record_index": ("Record index", "", "Sequential identifier of the parent exported record."),
    "item_index": ("Item index", "", "Sequential position in an exported one-dimensional array."),
    "parent_record_index": ("Parent record", "", "Record index linking this child row to its parent table."),
    "child_record_index": ("Child record", "", "Sequential child-row index within the parent record."),
    "group_index": ("Fit group", "", "Index of the fitted peak group or overlapping window."),
    "phase_index": ("Phase index", "", "Stable phase number within the refinement result."),
    "reflection_index": ("Reflection index", "", "Sequential reflection number within the exported phase/result."),
    "window_min": ("Fit window minimum", "° 2θ", "Lower bound of the fitted angular window."),
    "window_max": ("Fit window maximum", "° 2θ", "Upper bound of the fitted angular window."),
    "delta_two_theta_deg": ("Δ2θ", "°", "Observed minus reference or calculated peak position."),
    "score": ("Score", "", "Method-specific quality, match, or objective score."),
    "model": ("Model", "", "Selected mathematical or physical model."),
    "status": ("Status", "", "Analysis or validation status."),
    "warning": ("Warning", "", "Scientific or numerical warning."),
    "reason": ("Reason", "", "Recorded reason for a decision, exclusion, or validity state."),
}


@dataclass(frozen=True)
class ExportColumn:
    key: str
    label: str
    unit: str = ""
    description: str = ""
    number_format: str | None = None

    @property
    def heading(self) -> str:
        return f"{self.label} [{self.unit}]" if self.unit else self.label


@dataclass
class ExportTable:
    key: str
    title: str
    category: str
    description: str
    columns: list[ExportColumn]
    rows: list[dict[str, Any]]
    dataset_uid: str = ""
    dataset_name: str = ""
    warnings: list[str] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)

    def validate(self) -> None:
        keys = [column.key for column in self.columns]
        if len(keys) != len(set(keys)):
            raise ValueError(f"Duplicate column key in export table {self.key!r}.")
        if not self.columns:
            raise ValueError(f"Export table {self.key!r} has no columns.")
        for index, row in enumerate(self.rows, start=1):
            if not isinstance(row, Mapping):
                raise ValueError(f"Row {index} in {self.key!r} is not a mapping.")


@dataclass
class ExportPackage:
    project_name: str
    created_utc: str
    metadata: dict[str, Any]
    tables: list[ExportTable]

    def validate(self) -> None:
        keys: set[str] = set()
        for table in self.tables:
            table.validate()
            if table.key in keys:
                raise ValueError(f"Duplicate export table key: {table.key}")
            keys.add(table.key)


def utc_timestamp() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def _plain(value: Any) -> Any:
    if np is not None and isinstance(value, np.ndarray):
        return value.tolist()
    if np is not None and isinstance(value, np.generic):
        return value.item()
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, Mapping):
        return {str(k): _plain(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_plain(v) for v in value]
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def _is_sequence_value(value: Any) -> bool:
    if isinstance(value, (str, bytes, bytearray, Mapping)):
        return False
    if np is not None and isinstance(value, np.ndarray):
        return value.ndim == 1
    return isinstance(value, (list, tuple))


def _sequence_list(value: Any) -> list[Any]:
    if np is not None and isinstance(value, np.ndarray):
        return value.tolist()
    return list(value)


def _excel_value(value: Any) -> Any:
    value = _plain(value)
    if value is None:
        return None
    if isinstance(value, str):
        # Prevent Excel from interpreting imported notes, formulas, paths, or
        # identifiers as executable formulas when the workbook is opened.
        return "'" + value if value.startswith(("=", "+", "-", "@")) else value
    if isinstance(value, (int, float, bool)):
        return value
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def _humanize(key: str) -> str:
    text = re.sub(r"[_\-.]+", " ", str(key)).strip()
    text = re.sub(r"\s+", " ", text)
    replacements = {
        "rwp": "Rwp",
        "rp": "Rp",
        "rexp": "Rexp",
        "qpa": "QPA",
        "cif": "CIF",
        "uid": "UID",
        "fwhm": "FWHM",
        "hkl": "hkl",
        "chi square": "χ²",
        "two theta": "2θ",
    }
    lowered = text.lower()
    if lowered in replacements:
        return replacements[lowered]
    words = []
    for word in text.split():
        words.append(replacements.get(word.lower(), word.capitalize()))
    return " ".join(words)


def _column_for_key(key: str) -> ExportColumn:
    if key in COLUMN_LABELS:
        label, unit, description = COLUMN_LABELS[key]
    else:
        label, unit, description = _humanize(key), "", ""
        lowered = key.lower()
        if lowered.endswith("_deg") or "two_theta" in lowered or lowered.endswith("_degree"):
            unit = "°"
        elif lowered.endswith("_angstrom") or lowered.startswith("cell_a") or lowered.startswith("cell_b") or lowered.startswith("cell_c"):
            unit = "Å"
        elif lowered.endswith("_percent") or lowered.endswith("_pct"):
            unit = "%"
        elif "seconds" in lowered or lowered.endswith("_s"):
            unit = "s"
        elif "temperature" in lowered:
            unit = "K"
    if not description:
        description = f"Exported analysis field '{key}'."
    number_format = None
    if unit in {"°", "° 2θ", "Å"}:
        number_format = "0.000000"
    elif unit in {"%", "wt%"}:
        number_format = "0.000"
    elif unit:
        number_format = "0.000000"
    return ExportColumn(key=key, label=label, unit=unit, description=description, number_format=number_format)


def _deduplicate_columns(columns: Sequence[ExportColumn]) -> list[ExportColumn]:
    used: dict[str, int] = {}
    output: list[ExportColumn] = []
    for column in columns:
        heading = column.heading.casefold()
        count = used.get(heading, 0) + 1
        used[heading] = count
        if count == 1:
            output.append(column)
            continue
        output.append(ExportColumn(
            key=column.key,
            label=f"{column.label} ({count})",
            unit=column.unit,
            description=column.description,
            number_format=column.number_format,
        ))
    return output


def _ordered_keys(records: Sequence[Mapping[str, Any]], preferred: Sequence[str] = ()) -> list[str]:
    ordered: list[str] = []
    for key in preferred:
        if any(key in row for row in records):
            ordered.append(key)
    for row in records:
        for key in row:
            key = str(key)
            if key not in ordered:
                ordered.append(key)
    return ordered


def _records_table(
    *,
    key: str,
    title: str,
    category: str,
    description: str,
    records: Iterable[Mapping[str, Any]],
    dataset_uid: str = "",
    dataset_name: str = "",
    preferred: Sequence[str] = (),
) -> ExportTable | None:
    rows = [{str(k): _plain(v) for k, v in dict(row).items()} for row in records if isinstance(row, Mapping)]
    if not rows:
        return None
    keys = _ordered_keys(rows, preferred)
    columns = _deduplicate_columns([_column_for_key(column) for column in keys])
    return ExportTable(
        key=key,
        title=title,
        category=category,
        description=description,
        columns=columns,
        rows=rows,
        dataset_uid=dataset_uid,
        dataset_name=dataset_name,
    )


def _flatten_record(record: Mapping[str, Any], prefix: str = "") -> tuple[dict[str, Any], dict[str, list[Mapping[str, Any]]]]:
    """Flatten nested scalar mappings and return child record lists separately."""
    scalars: dict[str, Any] = {}
    children: dict[str, list[Mapping[str, Any]]] = {}
    for raw_key, raw_value in record.items():
        key = f"{prefix}.{raw_key}" if prefix else str(raw_key)
        value = _plain(raw_value)
        if isinstance(value, Mapping):
            nested_scalars, nested_children = _flatten_record(value, key)
            scalars.update(nested_scalars)
            children.update(nested_children)
        elif isinstance(value, list) and value and all(isinstance(item, Mapping) for item in value):
            scalars[f"{key}_row_count"] = len(value)
            children[key] = list(value)
        elif isinstance(value, list):
            scalars[key] = "; ".join(str(item) for item in value) if len(value) <= 20 else f"{len(value)} value(s)"
        else:
            scalars[key] = value
    return scalars, children


def _records_tables(
    *,
    key: str,
    title: str,
    category: str,
    description: str,
    records: Iterable[Mapping[str, Any]],
    dataset_uid: str = "",
    dataset_name: str = "",
    preferred: Sequence[str] = (),
) -> list[ExportTable]:
    source = [dict(row) for row in records if isinstance(row, Mapping)]
    if not source:
        return []
    parent_rows: list[dict[str, Any]] = []
    child_rows: dict[str, list[dict[str, Any]]] = {}
    for parent_index, record in enumerate(source, start=1):
        flat, children = _flatten_record(record)
        flat = {"record_index": parent_index, **flat}
        parent_rows.append(flat)
        for child_path, child_records in children.items():
            bucket = child_rows.setdefault(child_path, [])
            for child_index, child_record in enumerate(child_records, start=1):
                child_flat, grandchildren = _flatten_record(child_record)
                child_flat = {
                    "parent_record_index": parent_index,
                    "child_record_index": child_index,
                    **child_flat,
                }
                if grandchildren:
                    child_flat["nested_table_warning"] = (
                        "Additional nested record lists were summarized; inspect the source result if required."
                    )
                bucket.append(child_flat)
    tables: list[ExportTable] = []
    parent = _records_table(
        key=key, title=title, category=category, description=description,
        records=parent_rows, dataset_uid=dataset_uid, dataset_name=dataset_name,
        preferred=("record_index", *preferred),
    )
    if parent:
        tables.append(parent)
    for child_path, rows in child_rows.items():
        child = _records_table(
            key=f"{key}_{safe_filename(child_path)}",
            title=f"{title}: {_humanize(child_path)}",
            category=category,
            description=f"Child rows extracted from '{child_path}' so nested tables are not stored as unreadable JSON cells.",
            records=rows, dataset_uid=dataset_uid, dataset_name=dataset_name,
            preferred=("parent_record_index", "child_record_index", "phase_index", "center", "position", "h", "k", "l"),
        )
        if child:
            tables.append(child)
    return tables


def _flatten_scalars(value: Any, prefix: str = "") -> list[dict[str, Any]]:
    value = _plain(value)
    rows: list[dict[str, Any]] = []
    if isinstance(value, Mapping):
        for key in sorted(value, key=lambda x: str(x)):
            child = f"{prefix}.{key}" if prefix else str(key)
            item = value[key]
            if isinstance(item, Mapping):
                rows.extend(_flatten_scalars(item, child))
            elif isinstance(item, list) and item and all(isinstance(x, Mapping) for x in item):
                rows.append({"parameter": child, "value": f"{len(item)} row(s)", "data_type": "table"})
            elif isinstance(item, list):
                rows.append({"parameter": child, "value": f"{len(item)} value(s)", "data_type": "array"})
            else:
                rows.append({"parameter": child, "value": item, "data_type": type(item).__name__})
    else:
        rows.append({"parameter": prefix or "value", "value": value, "data_type": type(value).__name__})
    return rows


def _summary_table(
    *, key: str, title: str, category: str, description: str, value: Any,
    dataset_uid: str = "", dataset_name: str = ""
) -> ExportTable | None:
    rows = _flatten_scalars(value)
    if not rows:
        return None
    return ExportTable(
        key=key,
        title=title,
        category=category,
        description=description,
        columns=[
            ExportColumn("parameter", "Parameter", description="Stable dotted path to the exported value."),
            ExportColumn("value", "Value", description="Scalar value or concise table/array summary."),
            ExportColumn("data_type", "Data type", description="Original logical value type."),
        ],
        rows=rows,
        dataset_uid=dataset_uid,
        dataset_name=dataset_name,
    )


def _profile_table(dataset: Mapping[str, Any]) -> ExportTable:
    x = list(_plain(dataset.get("x", [])) or [])
    raw = list(_plain(dataset.get("y_raw", [])) or [])
    processed_value = dataset.get("y_processed")
    processed = list(_plain(processed_value) or []) if processed_value is not None else []
    if len(x) != len(raw):
        raise ValueError(f"Dataset {dataset.get('name')} has misaligned x/raw arrays.")
    if processed and len(processed) != len(x):
        raise ValueError(f"Dataset {dataset.get('name')} has misaligned processed intensity.")
    rows = []
    for index, xx in enumerate(x):
        row = {
            "point_index": index + 1,
            "two_theta_deg": xx,
            "intensity_raw": raw[index],
            "intensity_processed": processed[index] if processed else None,
            "processed_minus_raw": (processed[index] - raw[index]) if processed else None,
        }
        rows.append(row)
    return ExportTable(
        key=f"dataset_{safe_filename(dataset.get('uid') or dataset.get('name') or 'data')}_profile",
        title=f"Diffraction profile — {dataset.get('name', 'Dataset')}",
        category="pattern",
        description="Complete imported 2θ grid with raw and processed intensity kept in separate columns. Blank processed values mean no processed profile was stored.",
        columns=[
            _column_for_key("point_index"),
            _column_for_key("two_theta_deg"),
            _column_for_key("intensity_raw"),
            _column_for_key("intensity_processed"),
            ExportColumn("processed_minus_raw", "Processed − raw", "intensity", "Point-wise change introduced by preprocessing."),
        ],
        rows=rows,
        dataset_uid=str(dataset.get("uid", "")),
        dataset_name=str(dataset.get("name", "")),
    )


def _aligned_profile_from_result(result: Mapping[str, Any]) -> dict[str, Sequence[Any]]:
    aliases = [
        ("two_theta_deg", ("observed_x", "x", "two_theta", "two_theta_deg")),
        ("observed_intensity", ("observed_y", "y_observed", "observed_intensity")),
        ("calculated_intensity", ("calculated_y", "y_calculated", "calculated_intensity")),
        ("background_intensity", ("background_y", "y_background", "background_intensity")),
        ("difference_observed_minus_calculated", ("difference_y", "residual_y", "difference")),
        ("observed_sigma", ("observed_sigma", "sigma")),
        ("statistical_weight", ("statistical_weights", "weights")),
    ]
    output: dict[str, Sequence[Any]] = {}
    for target, candidates in aliases:
        for candidate in candidates:
            value = result.get(candidate)
            if isinstance(value, (list, tuple)) or (np is not None and isinstance(value, np.ndarray)):
                output[target] = list(_plain(value))
                break
    lengths_by_name = {name: len(values) for name, values in output.items()}
    if len(set(lengths_by_name.values())) > 1:
        detail = ", ".join(f"{name}={length}" for name, length in lengths_by_name.items())
        raise ValueError(f"Cannot export a misaligned refinement/profile result: {detail}")
    return output


def _clean_table_filename(dataset_name: str, order: int, suffix: str) -> str:
    """Return a readable, stable filename for the compact TXT package."""
    return f"{safe_filename(dataset_name or 'dataset')}_{order:02d}_{safe_filename(suffix)}.txt"


def _clean_dataset_filename(dataset: Mapping[str, Any], order: int, suffix: str) -> str:
    base = str(dataset.get("_clean_export_filebase") or dataset.get("name") or "dataset")
    return _clean_table_filename(base, order, suffix)


def _with_clean_filename(table: ExportTable, filename: str) -> ExportTable:
    table.metadata["clean_filename"] = filename
    return table


def _decoded_contract_result(
    analyses: Mapping[str, Any],
    kind: str,
    legacy_key: str,
) -> dict[str, Any]:
    """Merge a rich legacy result with its authoritative typed contract.

    Result contracts intentionally preserve the core numerical arrays.  The
    live result contains useful presentation fields such as the Pawley/Le Bail
    mode and detailed parameter rows.  Merging both gives the compact exporter
    authoritative profiles without discarding those human-readable details.
    """
    legacy = analyses.get(legacy_key)
    merged = dict(_plain(legacy)) if isinstance(legacy, Mapping) else {}
    contracts = analyses.get("result_contracts")
    envelope = contracts.get(kind) if isinstance(contracts, Mapping) else None
    if not isinstance(envelope, Mapping):
        return merged
    decoded = decode_value(envelope)
    if not isinstance(decoded, Mapping):
        return merged
    outputs = decoded.get("numerical_outputs")
    if isinstance(outputs, Mapping):
        outputs = dict(_plain(outputs))
        statistics = outputs.get("statistics")
        if isinstance(statistics, Mapping):
            for key, value in statistics.items():
                merged[str(key)] = value
        merged.update(outputs)
    parameters = decoded.get("parameters")
    if isinstance(parameters, Mapping) and parameters:
        merged.setdefault("settings", dict(_plain(parameters)))
    merged["contract_validation_status"] = decoded.get("validation_status")
    merged["contract_engine"] = decoded.get("engine")
    return merged


def _first_value(source: Mapping[str, Any], *keys: str) -> Any:
    for key in keys:
        value = source.get(key)
        if value is not None:
            return value
    return None


def _clean_raw_table(dataset: Mapping[str, Any]) -> ExportTable:
    x = list(_plain(dataset.get("x", [])) or [])
    raw = list(_plain(dataset.get("y_raw", [])) or [])
    if len(x) != len(raw):
        raise ValueError(f"Dataset {dataset.get('name')} has misaligned x/raw arrays.")
    name = str(dataset.get("name") or "Dataset")
    uid = str(dataset.get("uid") or "")
    return _with_clean_filename(ExportTable(
        key=f"{safe_filename(uid or name)}_raw_data",
        title=f"Diffraction profile — raw data — {name}",
        category="pattern",
        description="Imported measurement values, unchanged by preprocessing.",
        columns=[
            _column_for_key("point_index"),
            _column_for_key("two_theta_deg"),
            _column_for_key("intensity_raw"),
        ],
        rows=[
            {"point_index": index + 1, "two_theta_deg": xx, "intensity_raw": raw[index]}
            for index, xx in enumerate(x)
        ],
        dataset_uid=uid,
        dataset_name=name,
    ), _clean_dataset_filename(dataset, 1, "raw_data"))


def _clean_treated_table(dataset: Mapping[str, Any], analyses: Mapping[str, Any]) -> ExportTable | None:
    x = list(_plain(dataset.get("x", [])) or [])
    treated_value = dataset.get("y_processed")
    treated = list(_plain(treated_value) or []) if treated_value is not None else []
    profile_rows = analyses.get("preprocessing_profile")
    profile_rows = list(profile_rows) if isinstance(profile_rows, (list, tuple)) else []

    if not treated and profile_rows:
        treated = [row.get("intensity_processed") for row in profile_rows if isinstance(row, Mapping)]
    if not treated:
        contract = _decoded_contract_result(analyses, "preprocessing", "preprocessing")
        treated = list(_plain(_first_value(contract, "processed_profile", "processed")) or [])
    if not treated or len(treated) != len(x):
        return None

    background: list[Any] = []
    smoothed: list[Any] = []
    if len(profile_rows) == len(x):
        background = [
            _first_value(row, "background_raw_units", "background_display_units")
            if isinstance(row, Mapping) else None
            for row in profile_rows
        ]
        smoothed = [row.get("smoothed_intensity") if isinstance(row, Mapping) else None for row in profile_rows]
    if not background or not any(value is not None for value in background):
        contract = _decoded_contract_result(analyses, "preprocessing", "preprocessing")
        background = list(_plain(_first_value(contract, "background_profile", "background")) or [])
        smoothed = list(_plain(_first_value(contract, "smoothed_profile", "smoothed")) or [])

    include_background = len(background) == len(x) and any(value is not None for value in background)
    include_smoothed = (
        len(smoothed) == len(x)
        and any(value is not None for value in smoothed)
        and any(
            treated[index] is None
            or value is None
            or abs(float(value) - float(treated[index])) > 1e-12
            for index, value in enumerate(smoothed)
        )
    )
    rows = []
    for index, xx in enumerate(x):
        row = {
            "point_index": index + 1,
            "two_theta_deg": xx,
            "intensity_treated": treated[index],
        }
        if include_background:
            row["background_intensity"] = background[index]
        if include_smoothed:
            row["smoothed_intensity"] = smoothed[index]
        rows.append(row)
    columns = [
        _column_for_key("point_index"),
        _column_for_key("two_theta_deg"),
        _column_for_key("intensity_treated"),
    ]
    if include_background:
        columns.append(_column_for_key("background_intensity"))
    if include_smoothed:
        columns.append(ExportColumn(
            "smoothed_intensity", "Smoothed intensity", "processed units",
            "Intermediate smoothed profile when distinct from the final treated profile.",
        ))
    name = str(dataset.get("name") or "Dataset")
    uid = str(dataset.get("uid") or "")
    return _with_clean_filename(ExportTable(
        key=f"{safe_filename(uid or name)}_treated_data",
        title=f"Diffraction profile — treated data — {name}",
        category="preprocessing",
        description="Final analysis-ready profile with only the useful preprocessing columns.",
        columns=columns,
        rows=rows,
        dataset_uid=uid,
        dataset_name=name,
    ), _clean_dataset_filename(dataset, 2, "treated_data"))


def _clean_peak_table(dataset: Mapping[str, Any], analyses: Mapping[str, Any]) -> ExportTable | None:
    source = analyses.get("peaks")
    records = [dict(row) for row in source] if isinstance(source, (list, tuple)) else []
    if not records:
        contract = _decoded_contract_result(analyses, "peak_list", "peaks")
        candidate = contract.get("peaks", contract.get("rows", []))
        records = [dict(row) for row in candidate] if isinstance(candidate, (list, tuple)) else []
    if not records:
        return None

    rows: list[dict[str, Any]] = []
    for index, record in enumerate(records, start=1):
        rows.append({
            "peak_index": index,
            "included": _first_value(record, "use", "included"),
            "manual": record.get("manual"),
            "two_theta_deg": _first_value(record, "position", "position_deg", "two_theta_deg", "center"),
            "intensity": _first_value(record, "intensity", "height", "amplitude"),
            "prominence": _first_value(record, "prominence", "prominence_percent"),
            "fwhm_deg": _first_value(record, "fwhm", "fwhm_deg"),
            "area": record.get("area"),
            "model": record.get("model"),
            "uncertainty": _first_value(record, "uncertainty", "position_standard_error", "center_standard_error"),
            "origin": _first_value(record, "origin", "method"),
        })
    preferred = (
        "peak_index", "included", "manual", "two_theta_deg", "intensity",
        "prominence", "fwhm_deg", "area", "model", "uncertainty", "origin",
    )
    keys = [key for key in preferred if key == "peak_index" or any(row.get(key) not in (None, "") for row in rows)]
    name = str(dataset.get("name") or "Dataset")
    uid = str(dataset.get("uid") or "")
    return _with_clean_filename(ExportTable(
        key=f"{safe_filename(uid or name)}_peaks",
        title=f"Peak list — {name}",
        category="peaks",
        description="Curated detected and manual peaks in one concise table.",
        columns=[_column_for_key(key) for key in keys],
        rows=[{key: row.get(key) for key in keys} for row in rows],
        dataset_uid=uid,
        dataset_name=name,
    ), _clean_dataset_filename(dataset, 3, "peaks"))


def _refinement_method(result: Mapping[str, Any], *, rietveld: bool = False) -> tuple[str, str]:
    if rietveld:
        return "Rietveld", "rietveld"
    raw = str(_first_value(result, "mode", "method", "refinement_mode") or "Whole-pattern")
    lowered = raw.casefold().replace("-", " ")
    if "pawley" in lowered:
        return "Pawley", "pawley"
    if "bail" in lowered:
        return "Le Bail", "le_bail"
    return "Pawley / Le Bail", "whole_pattern"


def _clean_refinement_summary(
    dataset: Mapping[str, Any], result: Mapping[str, Any], method: str, slug: str, order: int,
) -> ExportTable:
    statistics = result.get("statistics") if isinstance(result.get("statistics"), Mapping) else {}
    profile = result.get("profile") if isinstance(result.get("profile"), Mapping) else {}
    candidates = [
        ("method", method, ""),
        ("success", result.get("success"), "yes/no"),
        ("validation_status", result.get("contract_validation_status"), ""),
        ("rwp_percent", _first_value(result, "rwp_percent") if result.get("rwp_percent") is not None else statistics.get("rwp_percent"), "%"),
        ("rp_percent", _first_value(result, "rp_percent") if result.get("rp_percent") is not None else statistics.get("rp_percent"), "%"),
        ("rexp_percent", _first_value(result, "rexp_percent") if result.get("rexp_percent") is not None else statistics.get("rexp_percent"), "%"),
        ("goodness_of_fit_sqrt", _first_value(result, "goodness_of_fit_sqrt", "goodness_of_fit") if _first_value(result, "goodness_of_fit_sqrt", "goodness_of_fit") is not None else statistics.get("goodness_of_fit"), ""),
        ("reduced_chi_square", _first_value(result, "reduced_chi_square") if result.get("reduced_chi_square") is not None else statistics.get("reduced_chi_square"), ""),
        ("r_squared", _first_value(result, "r_squared") if result.get("r_squared") is not None else statistics.get("r_squared"), ""),
        ("rmse", _first_value(result, "rmse") if result.get("rmse") is not None else statistics.get("rmse"), "intensity"),
        ("statistics_valid", result.get("statistics_valid", statistics.get("statistics_valid")), "yes/no"),
        ("weighting_model", _first_value(result, "weighting_model", "weighting"), ""),
        ("profile_model", profile.get("model", result.get("profile_model")), ""),
        ("zero_shift_deg", result.get("zero_shift_deg"), "° 2θ"),
        ("background_order", result.get("background_order"), ""),
        ("data_point_count", result.get("data_point_count"), "points"),
        ("reflection_count", result.get("reflection_count"), "reflections"),
        ("parameter_count", result.get("parameter_count"), "parameters"),
        ("elapsed_seconds", result.get("elapsed_seconds"), "s"),
        ("acceleration", result.get("acceleration"), ""),
        ("statistics_reason", result.get("statistics_reason", statistics.get("statistics_reason")), ""),
    ]
    rows = [
        {"parameter": key, "value": _plain(value), "unit": unit, "standard_uncertainty": None}
        for key, value, unit in candidates if value not in (None, "", [])
    ]
    parameters = result.get("parameters")
    if isinstance(parameters, (list, tuple)):
        for parameter in parameters:
            if not isinstance(parameter, Mapping):
                continue
            rows.append({
                "parameter": _first_value(parameter, "parameter", "name", "label"),
                "value": _first_value(parameter, "value", "refined_value"),
                "unit": parameter.get("unit", ""),
                "standard_uncertainty": _first_value(parameter, "standard_error", "uncertainty", "esd"),
            })
    name = str(dataset.get("name") or "Dataset")
    uid = str(dataset.get("uid") or "")
    return _with_clean_filename(ExportTable(
        key=f"{safe_filename(uid or name)}_{slug}_summary",
        title=f"{method} refinement summary and parameters — {name}",
        category="refinement",
        description="Fit quality, model settings and refined scalar parameters.",
        columns=[
            ExportColumn("parameter", "Parameter", description="Reported statistic, setting or refined parameter."),
            ExportColumn("value", "Value", description="Reported value."),
            ExportColumn("unit", "Unit", description="Physical unit where applicable."),
            ExportColumn("standard_uncertainty", "Standard uncertainty", description="Estimated standard uncertainty where available."),
        ],
        rows=rows,
        dataset_uid=uid,
        dataset_name=name,
    ), _clean_dataset_filename(dataset, order, f"{slug}_summary"))


def _clean_refinement_profile(
    dataset: Mapping[str, Any], result: Mapping[str, Any], method: str, slug: str, order: int,
) -> ExportTable | None:
    profile = _aligned_profile_from_result(result)
    required = {"two_theta_deg", "observed_intensity", "calculated_intensity"}
    if not required.issubset(profile):
        return None
    if "difference_observed_minus_calculated" not in profile:
        profile["difference_observed_minus_calculated"] = [
            float(observed) - float(calculated)
            for observed, calculated in zip(profile["observed_intensity"], profile["calculated_intensity"])
        ]
    ordered = [
        key for key in (
            "two_theta_deg", "observed_intensity", "calculated_intensity",
            "background_intensity", "difference_observed_minus_calculated",
            "observed_sigma", "statistical_weight",
        ) if key in profile
    ]
    count = len(profile["two_theta_deg"])
    name = str(dataset.get("name") or "Dataset")
    uid = str(dataset.get("uid") or "")
    return _with_clean_filename(ExportTable(
        key=f"{safe_filename(uid or name)}_{slug}_profile",
        title=f"{method} observed / calculated profile — {name}",
        category="refinement",
        description="Plot-ready refinement profile with observed, calculated, background and observed-minus-calculated values.",
        columns=[_column_for_key(key) for key in ordered],
        rows=[{key: profile[key][index] for key in ordered} for index in range(count)],
        dataset_uid=uid,
        dataset_name=name,
    ), _clean_dataset_filename(dataset, order, f"{slug}_profile"))


def _clean_reflections_table(
    dataset: Mapping[str, Any], result: Mapping[str, Any], method: str, slug: str, order: int,
) -> ExportTable | None:
    source = result.get("reflections")
    records = [dict(row) for row in source] if isinstance(source, (list, tuple)) else []
    if not records:
        return None
    rows = []
    for index, row in enumerate(records, start=1):
        hkl = row.get("hkl")
        if isinstance(hkl, (list, tuple)) and len(hkl) >= 3:
            h, k, l = hkl[:3]
        else:
            h, k, l = row.get("h"), row.get("k"), row.get("l")
        rows.append({
            "reflection_index": row.get("reflection_index", index),
            "phase_name": _first_value(row, "phase_name", "phase"),
            "h": h,
            "k": k,
            "l": l,
            "hkl_label": row.get("hkl_label"),
            "two_theta_deg": _first_value(row, "two_theta_deg", "position", "calculated_two_theta_deg"),
            "d_spacing_angstrom": _first_value(row, "d_spacing", "d_spacing_angstrom"),
            "fwhm_deg": _first_value(row, "fwhm_deg", "fwhm"),
            "intensity": _first_value(row, "extracted_intensity", "structure_intensity", "intensity", "reference_intensity"),
            "uncertainty": _first_value(row, "intensity_standard_error", "standard_error", "uncertainty"),
            "overlap_group": _first_value(row, "overlap_group_id", "overlap_group"),
            "status": _first_value(row, "intensity_identifiability", "status"),
            "included": _first_value(row, "included", "use"),
        })
    preferred = (
        "reflection_index", "phase_name", "h", "k", "l", "hkl_label", "two_theta_deg",
        "d_spacing_angstrom", "fwhm_deg", "intensity", "uncertainty", "overlap_group",
        "status", "included",
    )
    keys = [key for key in preferred if key == "reflection_index" or any(row.get(key) not in (None, "") for row in rows)]
    name = str(dataset.get("name") or "Dataset")
    uid = str(dataset.get("uid") or "")
    return _with_clean_filename(ExportTable(
        key=f"{safe_filename(uid or name)}_{slug}_reflections",
        title=f"{method} reflections — {name}",
        category="refinement",
        description="Concise reflection list used by the refinement.",
        columns=[_column_for_key(key) for key in keys],
        rows=[{key: row.get(key) for key in keys} for row in rows],
        dataset_uid=uid,
        dataset_name=name,
    ), _clean_dataset_filename(dataset, order, f"{slug}_reflections"))


def _clean_phases_table(
    dataset: Mapping[str, Any], result: Mapping[str, Any], order: int,
) -> ExportTable | None:
    source = result.get("phases")
    phases = [dict(row) for row in source] if isinstance(source, (list, tuple)) else []
    if not phases:
        return None
    rows = []
    for index, phase in enumerate(phases, start=1):
        cell = phase.get("refined_cell") if isinstance(phase.get("refined_cell"), Mapping) else {}
        rows.append({
            "phase_index": phase.get("phase_index", index),
            "phase_name": _first_value(phase, "phase_name", "name", "data_name"),
            "formula": phase.get("formula"),
            "crystal_system": phase.get("crystal_system"),
            "scale_factor": phase.get("scale_factor"),
            "pattern_fraction_percent": _first_value(phase, "weight_fraction_percent", "pattern_scale_fraction_percent", "pattern_fraction_percent", "phase_fraction_percent"),
            "cell_a_angstrom": cell.get("a"),
            "cell_b_angstrom": cell.get("b"),
            "cell_c_angstrom": cell.get("c"),
            "cell_alpha_deg": cell.get("alpha"),
            "cell_beta_deg": cell.get("beta"),
            "cell_gamma_deg": cell.get("gamma"),
            "reflection_count": phase.get("reflection_count"),
        })
    keys = [key for key in rows[0] if key == "phase_index" or any(row.get(key) not in (None, "") for row in rows)]
    name = str(dataset.get("name") or "Dataset")
    uid = str(dataset.get("uid") or "")
    return _with_clean_filename(ExportTable(
        key=f"{safe_filename(uid or name)}_rietveld_phases",
        title=f"Rietveld phase results — {name}",
        category="refinement",
        description="Phase scales, fractions and refined unit-cell values.",
        columns=[_column_for_key(key) for key in keys],
        rows=[{key: row.get(key) for key in keys} for row in rows],
        dataset_uid=uid,
        dataset_name=name,
    ), _clean_dataset_filename(dataset, order, "rietveld_phases"))


def build_clean_export_package(
    snapshot: Mapping[str, Any],
    *,
    include_content: set[str] | None = None,
) -> ExportPackage:
    """Build the concise, user-facing scientific export package.

    Raw, treated, peak and refinement products are intentionally separated so
    every TXT file has one clear purpose and can be plotted without cleanup.
    """
    include = set(include_content or CLEAN_CONTENT_LABELS)
    datasets = [row for row in snapshot.get("datasets", []) if isinstance(row, Mapping)]
    name_totals: dict[str, int] = {}
    for dataset in datasets:
        base = safe_filename(str(dataset.get("name") or "Dataset"))
        name_totals[base.casefold()] = name_totals.get(base.casefold(), 0) + 1
    tables: list[ExportTable] = []
    duplicate_indices: dict[str, int] = {}
    for source_dataset in datasets:
        dataset = dict(source_dataset)
        base = safe_filename(str(dataset.get("name") or "Dataset"))
        folded = base.casefold()
        if name_totals.get(folded, 0) > 1:
            duplicate_indices[folded] = duplicate_indices.get(folded, 0) + 1
            uid = safe_filename(str(dataset.get("uid") or ""))[:8]
            qualifier = uid or str(duplicate_indices[folded])
            dataset["_clean_export_filebase"] = f"{base}_{qualifier}"
        analyses = dataset.get("analyses") if isinstance(dataset.get("analyses"), Mapping) else {}
        if "raw" in include:
            tables.append(_clean_raw_table(dataset))
        if "treated" in include:
            treated = _clean_treated_table(dataset, analyses)
            if treated is not None:
                tables.append(treated)
        if "peaks" in include:
            peaks = _clean_peak_table(dataset, analyses)
            if peaks is not None:
                tables.append(peaks)
        if "refinement" in include:
            whole = _decoded_contract_result(analyses, "whole_pattern_refinement", "whole_pattern")
            if whole:
                method, slug = _refinement_method(whole)
                tables.append(_clean_refinement_summary(dataset, whole, method, slug, 4))
                for candidate in (
                    _clean_refinement_profile(dataset, whole, method, slug, 4),
                    _clean_reflections_table(dataset, whole, method, slug, 4),
                ):
                    if candidate is not None:
                        tables.append(candidate)
            rietveld = _decoded_contract_result(analyses, "rietveld_refinement", "rietveld")
            if rietveld:
                method, slug = _refinement_method(rietveld, rietveld=True)
                tables.append(_clean_refinement_summary(dataset, rietveld, method, slug, 5))
                for candidate in (
                    _clean_refinement_profile(dataset, rietveld, method, slug, 5),
                    _clean_phases_table(dataset, rietveld, 5),
                    _clean_reflections_table(dataset, rietveld, method, slug, 5),
                ):
                    if candidate is not None:
                        tables.append(candidate)

    project_name = str(snapshot.get("project_name") or "Afruz PXRD Project")
    package = ExportPackage(
        project_name=project_name,
        created_utc=utc_timestamp(),
        metadata={
            "project_name": project_name,
            "project_path": snapshot.get("project_path") or "Not saved",
            "application_version": snapshot.get("application_version") or APP_VERSION,
            "export_schema_version": EXPORT_SCHEMA_VERSION,
            "export_style": "clean_scientific",
            "created_utc": utc_timestamp(),
            "dataset_count": len(datasets),
            "table_count": len(tables),
        },
        tables=tables,
    )
    package.validate()
    return package


def _result_tables(
    *,
    dataset: Mapping[str, Any],
    result_key: str,
    title: str,
    category: str,
    result: Any,
    _depth: int = 0,
) -> list[ExportTable]:
    if not isinstance(result, Mapping) or not result:
        return []
    uid = str(dataset.get("uid", ""))
    name = str(dataset.get("name", ""))
    prefix = f"{safe_filename(uid or name)}_{safe_filename(result_key)}"
    tables: list[ExportTable] = []

    summary = _summary_table(
        key=f"{prefix}_summary",
        title=f"{title} summary — {name}",
        category=category,
        description="Scalar parameters, fit statistics, settings, warnings, and provenance. Large arrays and record lists are exported in their own tables.",
        value=result,
        dataset_uid=uid,
        dataset_name=name,
    )
    if summary:
        tables.append(summary)

    profile = _aligned_profile_from_result(result)
    if profile:
        keys = list(profile)
        n = len(next(iter(profile.values())))
        rows = [{key: profile[key][i] for key in keys} for i in range(n)]
        tables.append(ExportTable(
            key=f"{prefix}_profile",
            title=f"{title} profile — {name}",
            category=category,
            description="Aligned numerical profile. Observed, calculated, background, residual, sigma, and weight columns remain row-for-row compatible.",
            columns=[_column_for_key(key) for key in keys],
            rows=rows,
            dataset_uid=uid,
            dataset_name=name,
        ))

    if result_key in {"rietveld", "multicomponent", "whole_pattern"} and result.get("reflections"):
        manuscript_rows, manuscript_metadata = build_manuscript_plot_data(result)
        if manuscript_rows:
            tables.append(ExportTable(
                key=f"{prefix}_plot_data_for_manuscript",
                title=f"Manuscript plot data — {name}",
                category=category,
                description=(
                    "Single long-format plot table. Filter Record type = PROFILE for observed, calculated, "
                    "background and difference curves; filter Record type = BRAGG for phase-resolved Bragg lines. "
                    "Rwp, Rp, Rexp, Rwp/Rexp (GoF), and reduced chi-square are included. Display-offset columns "
                    "are plotting conveniences only and do not alter the scientific refinement result."
                ),
                columns=[ExportColumn(key, label, unit, description) for key, label, unit, description in MANUSCRIPT_COLUMN_SPECS],
                rows=manuscript_rows,
                dataset_uid=uid,
                dataset_name=name,
                metadata=manuscript_metadata,
            ))

    handled = {
        "observed_x", "x", "two_theta", "two_theta_deg", "observed_y", "y_observed",
        "observed_intensity", "calculated_y", "y_calculated", "calculated_intensity",
        "background_y", "y_background", "background_intensity", "difference_y",
        "residual_y", "difference", "observed_sigma", "sigma", "statistical_weights", "weights",
    }
    preferred_lists = {
        "phases": "Phase results",
        "reflections": "Reflection results",
        "peaks": "Peak results",
        "matches": "Matched peaks",
        "components": "Fit components",
        "parameters": "Refined parameters",
        "mismatches": "Mismatch review",
        "exclusions": "Excluded windows",
        "history": "Refinement history",
        "stages": "Refinement stages",
        "warnings": "Warnings",
        "audit_results": "Audit results",
        "robustness_results": "Robustness results",
    }
    scalar_sequences: dict[int, dict[str, list[Any]]] = {}
    for child_key, child in result.items():
        if child_key in handled:
            continue
        if _is_sequence_value(child):
            sequence = _sequence_list(child)
            if sequence and all(isinstance(item, Mapping) for item in sequence):
                tables.extend(_records_tables(
                    key=f"{prefix}_{safe_filename(child_key)}",
                    title=f"{title}: {preferred_lists.get(child_key, _humanize(child_key))} — {name}",
                    category=category,
                    description=f"Structured rows exported from result field '{child_key}'. Nested child records are exported as linked tables.",
                    records=sequence,
                    dataset_uid=uid,
                    dataset_name=name,
                    preferred=("phase_index", "phase_name", "hkl_label", "h", "k", "l", "two_theta_deg", "position", "included"),
                ))
            elif sequence:
                scalar_sequences.setdefault(len(sequence), {})[str(child_key)] = sequence
            continue
        if isinstance(child, Mapping) and child:
            if _depth < 3:
                tables.extend(_result_tables(
                    dataset=dataset,
                    result_key=f"{result_key}_{child_key}",
                    title=f"{title}: {_humanize(child_key)}",
                    category=category,
                    result=child,
                    _depth=_depth + 1,
                ))
            else:
                nested = _summary_table(
                    key=f"{prefix}_{safe_filename(child_key)}",
                    title=f"{title}: {_humanize(child_key)} — {name}",
                    category=category,
                    description=f"Nested result record '{child_key}' represented as parameter/value rows.",
                    value=child,
                    dataset_uid=uid,
                    dataset_name=name,
                )
                if nested:
                    tables.append(nested)

    for length, sequence_columns in sorted(scalar_sequences.items()):
        keys = list(sequence_columns)
        rows = []
        for item_index in range(length):
            row = {"item_index": item_index + 1}
            for key in keys:
                row[key] = sequence_columns[key][item_index]
            rows.append(row)
        suffix = "arrays" if len(sequence_columns) > 1 else safe_filename(keys[0])
        tables.append(ExportTable(
            key=f"{prefix}_{suffix}_{length}",
            title=f"{title}: {_humanize(' / '.join(keys))} — {name}",
            category=category,
            description=(
                "Aligned one-dimensional result values exported without truncation. "
                "Fields with the same length share a row index."
            ),
            columns=_deduplicate_columns([
                ExportColumn("item_index", "Item index", description="Sequential position in the source array."),
                *[_column_for_key(key) for key in keys],
            ]),
            rows=rows,
            dataset_uid=uid,
            dataset_name=name,
        ))
    return tables


def build_export_package(
    snapshot: Mapping[str, Any],
    *,
    include_categories: set[str] | None = None,
) -> ExportPackage:
    include = set(include_categories or CATEGORY_LABELS)
    project_name = str(snapshot.get("project_name") or "Afruz PXRD Project")
    datasets = [d for d in snapshot.get("datasets", []) if isinstance(d, Mapping)]
    tables: list[ExportTable] = []

    summary_rows = []
    for dataset in datasets:
        x = list(_plain(dataset.get("x", [])) or [])
        raw = list(_plain(dataset.get("y_raw", [])) or [])
        summary_rows.append({
            "dataset_name": dataset.get("name"),
            "dataset_uid": dataset.get("uid"),
            "source_path": dataset.get("source_path"),
            "point_count": len(x),
            "two_theta_min_deg": min(x) if x else None,
            "two_theta_max_deg": max(x) if x else None,
            "raw_intensity_min": min(raw) if raw else None,
            "raw_intensity_max": max(raw) if raw else None,
            "processed_profile_available": dataset.get("y_processed") is not None,
        })
    if "project" in include and summary_rows:
        table = _records_table(
            key="project_dataset_inventory",
            title="Project dataset inventory",
            category="project",
            description="One row per exported dataset with source, size, angular range, and processing availability.",
            records=summary_rows,
            preferred=("dataset_name", "dataset_uid", "source_path", "point_count", "two_theta_min_deg", "two_theta_max_deg"),
        )
        if table:
            tables.append(table)

    if "project" in include:
        project_summary = _summary_table(
            key="project_metadata",
            title="Project and export metadata",
            category="project",
            description="Project path, application version, export schema, active workflow state, and other reproducibility metadata.",
            value={**dict(snapshot.get("metadata", {})), "export_schema_version": EXPORT_SCHEMA_VERSION},
        )
        if project_summary:
            tables.append(project_summary)

    for dataset in datasets:
        uid = str(dataset.get("uid", ""))
        name = str(dataset.get("name", "Dataset"))
        analyses = dataset.get("analyses", {}) if isinstance(dataset.get("analyses"), Mapping) else {}
        result_contracts = (
            analyses.get("result_contracts", {})
            if isinstance(analyses.get("result_contracts"), Mapping)
            else {}
        )

        if "pattern" in include:
            tables.append(_profile_table(dataset))
            metadata = _summary_table(
                key=f"{safe_filename(uid or name)}_dataset_metadata",
                title=f"Dataset metadata — {name}",
                category="pattern",
                description="Imported file metadata and data provenance. Nested fields use stable dotted parameter paths.",
                value=dataset.get("metadata", {}),
                dataset_uid=uid,
                dataset_name=name,
            )
            if metadata:
                tables.append(metadata)

        contract_sections = [
            ("preprocessing", "Preprocessing result contract", "preprocessing"),
            ("peak_list", "Peak-list result contract", "peaks"),
            ("peak_fitting", "Peak-fitting result contract", "peaks"),
            ("instrument_calibration", "Instrument-calibration result contract", "validation"),
            ("unit_cell_refinement", "Unit-cell refinement result contract", "structure"),
            ("phase_identification", "Phase-identification result contract", "structure"),
            ("whole_pattern_refinement", "Whole-pattern refinement result contract", "refinement"),
            ("rietveld_refinement", "Rietveld refinement result contract", "refinement"),
            ("qpa", "Quantitative phase-analysis result contract", "qpa"),
            ("validation", "Validation result contract", "validation"),
        ]
        for result_kind, title, category in contract_sections:
            if category not in include:
                continue
            value = result_contracts.get(result_kind)
            if not isinstance(value, Mapping) or not value:
                continue
            decoded_value = decode_value(value)
            tables.extend(_result_tables(
                dataset=dataset,
                result_key=f"contract_{result_kind}",
                title=title,
                category=category,
                result=decoded_value,
            ))

        contract_legacy_keys = {
            "preprocessing": ("preprocessing",),
            "peak_list": ("peaks", "peak_list_metadata"),
            "peak_fitting": ("peak_fits", "fit_candidates"),
            "unit_cell_refinement": ("cell_refinement",),
            "phase_identification": ("phase_identification",),
            "whole_pattern_refinement": ("whole_pattern",),
            "rietveld_refinement": ("rietveld",),
            "qpa": ("exploratory_qpa", "validated_qpa"),
            "validation": ("validation",),
        }
        superseded_legacy_keys = {
            key
            for kind, keys in contract_legacy_keys.items()
            if kind in result_contracts
            for key in keys
        }
        section_map = [
            ("preprocessing", "Preprocessing", "preprocessing"),
            ("preprocessing_profile", "Preprocessing numerical profile", "preprocessing"),
            ("peaks", "Curated peaks", "peaks"),
            ("peak_list_metadata", "Peak-list metadata", "peaks"),
            ("peak_fits", "Peak fitting", "peaks"),
            ("fit_candidates", "Peak model candidates", "peaks"),
            ("size_strain", "Size and strain", "peaks"),
            ("cell_refinement", "Cell refinement", "structure"),
            ("phase_identification", "Phase identification", "structure"),
            ("phase_revolution", "Unknown-phase indexing", "structure"),
            ("structure_solution", "Structure solution", "structure"),
            ("whole_pattern", "Pawley / Le Bail refinement", "refinement"),
            ("rietveld", "Rietveld refinement", "refinement"),
            ("multicomponent", "Multiphase refinement", "refinement"),
            ("exploratory_qpa", "Exploratory QPA", "qpa"),
            ("validated_qpa", "Validated QPA", "qpa"),
            ("validation", "Validation and robustness", "validation"),
        ]
        for analysis_key, title, category in section_map:
            if category not in include:
                continue
            if analysis_key in superseded_legacy_keys:
                continue
            value = analyses.get(analysis_key)
            if value is None or value == {} or value == []:
                continue
            if isinstance(value, list) and value and all(isinstance(row, Mapping) for row in value):
                tables.extend(_records_tables(
                    key=f"{safe_filename(uid or name)}_{safe_filename(analysis_key)}",
                    title=f"{title} — {name}",
                    category=category,
                    description=f"Structured {title.lower()} rows for the selected dataset. Nested child records are exported as linked tables.",
                    records=value,
                    dataset_uid=uid,
                    dataset_name=name,
                ))
            else:
                tables.extend(_result_tables(
                    dataset=dataset,
                    result_key=analysis_key,
                    title=title,
                    category=category,
                    result=value if isinstance(value, Mapping) else {"values": value},
                ))

    project_analyses = snapshot.get("project_analyses", {})
    if isinstance(project_analyses, Mapping):
        project_sections = [
            ("cif_reference", "Project CIF reference", "structure"),
            ("residual_stress", "Residual stress", "validation"),
            ("instrument_profile", "Instrument profile", "validation"),
        ]
        for analysis_key, title, category in project_sections:
            if category not in include:
                continue
            value = project_analyses.get(analysis_key)
            if value is None or value == {} or value == []:
                continue
            tables.extend(_result_tables(
                dataset={"uid": "project", "name": "Project"},
                result_key=analysis_key,
                title=title,
                category=category,
                result=value if isinstance(value, Mapping) else {"values": value},
            ))

    batch_result = snapshot.get("batch_result")
    if "batch" in include and isinstance(batch_result, Mapping) and batch_result:
        tables.extend(_result_tables(
            dataset={"uid": "batch", "name": "Batch analysis"},
            result_key="batch",
            title="Batch analysis",
            category="batch",
            result=batch_result,
        ))

    package = ExportPackage(
        project_name=project_name,
        created_utc=utc_timestamp(),
        metadata={
            "project_name": project_name,
            "project_path": snapshot.get("project_path") or "Not saved",
            "application_version": snapshot.get("application_version") or APP_VERSION,
            "export_schema_version": EXPORT_SCHEMA_VERSION,
            "created_utc": utc_timestamp(),
            "dataset_count": len(datasets),
            "table_count": len(tables),
            **dict(snapshot.get("metadata", {})),
        },
        tables=tables,
    )
    package.validate()
    return package


def preview_package(package: ExportPackage) -> list[dict[str, Any]]:
    package.validate()
    rows = []
    for table in package.tables:
        rows.append({
            "include": True,
            "table_key": table.key,
            "table_title": table.title,
            "category": CATEGORY_LABELS.get(table.category, _humanize(table.category)),
            "dataset": table.dataset_name or "Project",
            "rows": len(table.rows),
            "columns": len(table.columns),
            "status": "Ready" if table.rows else "Header only",
            "warnings": "; ".join(table.warnings),
        })
    return rows


def _backup_existing(path: Path) -> Path | None:
    if not path.exists():
        return None
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    candidate = path.with_name(f"{path.stem}.backup_{stamp}{path.suffix}")
    counter = 1
    while candidate.exists():
        candidate = path.with_name(f"{path.stem}.backup_{stamp}_{counter}{path.suffix}")
        counter += 1
    shutil.copy2(path, candidate)
    return candidate


def _backup_existing_directory(path: Path) -> Path | None:
    if not path.exists():
        return None
    if not path.is_dir():
        raise ValueError(f"Expected an export directory but found a file: {path}")
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    candidate = path.with_name(f"{path.name}.backup_{stamp}")
    counter = 1
    while candidate.exists():
        candidate = path.with_name(f"{path.name}.backup_{stamp}_{counter}")
        counter += 1
    path.rename(candidate)
    return candidate


def _unique_sheet_name(title: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]+", " ", title).strip() or "Table"
    cleaned = re.sub(r"\s+", " ", cleaned)[:31]
    base = cleaned
    counter = 2
    while cleaned.lower() in used:
        suffix = f" {counter}"
        cleaned = (base[: 31 - len(suffix)] + suffix).strip()
        counter += 1
    used.add(cleaned.lower())
    return cleaned


def _set_widths(ws, min_row: int, max_row: int, max_col: int) -> None:
    for column_index in range(1, max_col + 1):
        width = 10
        for row_index in range(min_row, min(max_row, min_row + 250) + 1):
            value = ws.cell(row_index, column_index).value
            if value is None:
                continue
            width = max(width, min(42, len(str(value)) + 2))
        ws.column_dimensions[get_column_letter(column_index)].width = width


def write_excel_workbook(
    path: str | Path,
    package: ExportPackage,
    *,
    progress: Callable[[int, str], None] | None = None,
    cancelled: Callable[[], bool] | None = None,
) -> Path:
    package.validate()
    target = Path(path).with_suffix(".xlsx")
    target.parent.mkdir(parents=True, exist_ok=True)
    _backup_existing(target)

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    title_fill = PatternFill("solid", fgColor="17365D")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    section_fill = PatternFill("solid", fgColor="E2F0D9")
    title_font = Font(color="FFFFFF", bold=True, size=14)
    header_font = Font(color="17365D", bold=True)
    thin = Side(style="thin", color="B7C9D6")

    overview = wb.create_sheet(_unique_sheet_name("Overview", used))
    overview.sheet_view.showGridLines = False
    overview["A1"] = "Afruz PXRD Analyzer — Raptor Export Package"
    overview["A1"].fill = title_fill
    overview["A1"].font = title_font
    overview.merge_cells("A1:F1")
    row = 3
    for key, value in package.metadata.items():
        overview.cell(row, 1, _humanize(str(key)))
        overview.cell(row, 2, _excel_value(value))
        overview.cell(row, 1).font = header_font
        row += 1
    row += 1
    inventory_header_row = row
    inventory_headers = ["Table", "Category", "Dataset", "Rows", "Columns", "Worksheet", "Description"]
    for col, heading in enumerate(inventory_headers, start=1):
        cell = overview.cell(row, col, heading)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = Border(bottom=thin)
    row += 1

    sheet_records: list[tuple[ExportTable, str]] = []
    total = max(len(package.tables), 1)
    for index, table in enumerate(package.tables, start=1):
        if cancelled and cancelled():
            raise RuntimeError("Export cancelled by user.")
        sheet_name = _unique_sheet_name(table.title, used)
        sheet_records.append((table, sheet_name))
        ws = wb.create_sheet(sheet_name)
        ws.sheet_view.showGridLines = False
        max_col = max(1, len(table.columns))
        ws.cell(1, 1, table.title)
        ws.cell(1, 1).fill = title_fill
        ws.cell(1, 1).font = title_font
        if max_col > 1:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
        ws.cell(2, 1, table.description)
        ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
        if max_col > 1:
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
        ws.cell(3, 1, "Dataset")
        ws.cell(3, 2, table.dataset_name or "Project")
        ws.cell(3, 3, "Table key")
        ws.cell(3, 4, table.key)
        header_row = 5
        for col_index, column in enumerate(table.columns, start=1):
            cell = ws.cell(header_row, col_index, column.heading)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = Border(bottom=thin)
        for row_index, record in enumerate(table.rows, start=header_row + 1):
            for col_index, column in enumerate(table.columns, start=1):
                cell = ws.cell(row_index, col_index, _excel_value(record.get(column.key)))
                if column.number_format and isinstance(cell.value, (int, float)):
                    cell.number_format = column.number_format
                cell.alignment = Alignment(vertical="top")
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(max_col)}{max(header_row, header_row + len(table.rows))}"
        if table.rows:
            table_ref = f"A{header_row}:{get_column_letter(max_col)}{header_row + len(table.rows)}"
            table_name = re.sub(r"[^A-Za-z0-9_]", "_", f"T_{index}_{table.key}")[:250]
            if not table_name[0].isalpha():
                table_name = "T_" + table_name
            excel_table = Table(displayName=table_name, ref=table_ref)
            excel_table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
                showRowStripes=True, showColumnStripes=False,
            )
            ws.add_table(excel_table)
        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 36
        _set_widths(ws, 5, 5 + len(table.rows), max_col)
        if progress:
            progress(int(index * 75 / total), f"Writing Excel table {index}/{total}: {table.title}")

    for table, sheet_name in sheet_records:
        overview.cell(row, 1, table.title)
        overview.cell(row, 2, CATEGORY_LABELS.get(table.category, _humanize(table.category)))
        overview.cell(row, 3, table.dataset_name or "Project")
        overview.cell(row, 4, len(table.rows))
        overview.cell(row, 5, len(table.columns))
        overview.cell(row, 6, sheet_name)
        overview.cell(row, 7, table.description)
        overview.cell(row, 6).hyperlink = f"#'{sheet_name}'!A1"
        overview.cell(row, 6).style = "Hyperlink"
        overview.cell(row, 7).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    if sheet_records:
        inv_ref = f"A{inventory_header_row}:G{row - 1}"
        inv = Table(displayName="ExportTableInventory", ref=inv_ref)
        inv.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
        overview.add_table(inv)
    overview.freeze_panes = f"A{inventory_header_row + 1}"
    _set_widths(overview, 1, row, 7)
    overview.column_dimensions["G"].width = 56

    dictionary = wb.create_sheet(_unique_sheet_name("Data Dictionary", used))
    dictionary.sheet_view.showGridLines = False
    dictionary["A1"] = "Data dictionary"
    dictionary["A1"].fill = title_fill
    dictionary["A1"].font = title_font
    dictionary.merge_cells("A1:H1")
    dict_headers = ["Table key", "Table title", "Category", "Dataset", "Column key", "Column heading", "Unit", "Meaning"]
    for c, heading in enumerate(dict_headers, start=1):
        dictionary.cell(3, c, heading).fill = section_fill
        dictionary.cell(3, c).font = header_font
    dr = 4
    for table in package.tables:
        for column in table.columns:
            values = [
                table.key, table.title, CATEGORY_LABELS.get(table.category, table.category),
                table.dataset_name or "Project", column.key, column.heading, column.unit, column.description,
            ]
            for c, value in enumerate(values, start=1):
                dictionary.cell(dr, c, value)
            dr += 1
    if dr > 4:
        dt = Table(displayName="ExportDataDictionary", ref=f"A3:H{dr - 1}")
        dt.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
        dictionary.add_table(dt)
    dictionary.freeze_panes = "A4"
    _set_widths(dictionary, 1, dr, 8)
    dictionary.column_dimensions["H"].width = 54

    descriptor, temp_name = tempfile.mkstemp(prefix=f".{target.name}.", suffix=".tmp.xlsx", dir=str(target.parent))
    os.close(descriptor)
    temporary = Path(temp_name)
    try:
        if progress:
            progress(85, "Finalizing Excel workbook")
        wb.save(temporary)
        os.replace(temporary, target)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    if progress:
        progress(100, "Excel workbook complete")
    return target


def _readme_text(package: ExportPackage) -> str:
    lines = [
        "AFRUZ PXRD ANALYZER — RAPTOR EXPORT PACKAGE",
        "=" * 48,
        "",
        f"Project: {package.project_name}",
        f"Created UTC: {package.created_utc}",
        f"Export schema: {EXPORT_SCHEMA_VERSION}",
        f"Tables: {len(package.tables)}",
        "",
        "Format rules",
        "------------",
        "All numerical tables are UTF-8, TAB-delimited TXT files with LF line endings.",
        "Missing values are written as NA. Units are included in column headings.",
        "Every table is rectangular and has a stable table key. The manifest records SHA-256 hashes.",
        "Refinement exports additionally include plot-data-for-manuscrit_<dataset>.txt as one long-format, plot-ready table.",
        "",
        "Table inventory",
        "---------------",
    ]
    for table in package.tables:
        lines.append(f"- {table.key}: {table.title} ({len(table.rows)} rows × {len(table.columns)} columns)")
        lines.append(f"  {table.description}")
    return "\n".join(lines) + "\n"


def _clean_readme_text(package: ExportPackage) -> str:
    lines = [
        "AFRUZ PXRD ANALYZER — CLEAN SCIENTIFIC DATA",
        "=" * 48,
        "",
        f"Project: {package.project_name}",
        f"Created UTC: {package.created_utc}",
        "",
        "Each UTF-8 TXT file contains one TAB-delimited rectangular table.",
        "Units are written in the column headings and missing values are NA.",
        "Refinement profile difference is observed intensity minus calculated intensity.",
        "The manifest lists every data file and its SHA-256 checksum.",
        "",
        "Files",
        "-----",
    ]
    for table in package.tables:
        filename = str(table.metadata.get("clean_filename") or f"{safe_filename(table.key)}.txt")
        lines.append(f"- {filename}: {table.title}")
    return "\n".join(lines) + "\n"


def write_text_package(
    directory: str | Path,
    package: ExportPackage,
    *,
    progress: Callable[[int, str], None] | None = None,
    cancelled: Callable[[], bool] | None = None,
    compact: bool = False,
) -> Path:
    package.validate()
    root = Path(directory)
    root.mkdir(parents=True, exist_ok=True)
    files: list[Path] = []
    total = max(len(package.tables), 1)
    dictionary_rows: list[dict[str, Any]] = []
    inventory_rows: list[dict[str, Any]] = []
    for index, table in enumerate(package.tables, start=1):
        if cancelled and cancelled():
            raise RuntimeError("Export cancelled by user.")
        clean_filename = str(table.metadata.get("clean_filename") or "").strip()
        if compact and clean_filename:
            # The compact builder only creates a leaf filename.  Normalising it
            # again prevents a metadata value from escaping the export folder.
            filename = safe_filename(Path(clean_filename).name, "data.txt")
            if not filename.lower().endswith(".txt"):
                filename += ".txt"
        elif table.key.endswith("_plot_data_for_manuscript"):
            filename = f"plot-data-for-manuscrit_{safe_filename(table.dataset_name or 'project')}.txt"
            candidate = root / filename
            if candidate.exists():
                filename = f"plot-data-for-manuscrit_{safe_filename(table.dataset_name or 'project')}_{index:03d}.txt"
        else:
            filename = f"{index:03d}_{safe_filename(table.key)}.txt"
        headings = [column.heading for column in table.columns]
        heading_by_key = {column.key: column.heading for column in table.columns}
        rows = [{heading_by_key[column.key]: row.get(column.key) for column in table.columns} for row in table.rows]
        public_metadata = {
            key: value for key, value in table.metadata.items()
            if key != "clean_filename"
        }
        result = write_table_txt(
            root / filename,
            rows,
            columns=headings,
            title=table.title,
            metadata={
                "table_key": table.key,
                "category": CATEGORY_LABELS.get(table.category, table.category),
                "dataset_name": table.dataset_name or "Project",
                "dataset_uid": table.dataset_uid or "Not applicable",
                "description": table.description,
                "export_schema_version": EXPORT_SCHEMA_VERSION,
                **public_metadata,
            },
            backup=False,
        )
        table_path = Path(result["txt_path"])
        files.append(table_path)
        inventory_rows.append({
            "table_key": table.key,
            "table_title": table.title,
            "category": CATEGORY_LABELS.get(table.category, table.category),
            "dataset": table.dataset_name or "Project",
            "file": table_path.name,
            "rows": len(table.rows),
            "columns": len(table.columns),
            "description": table.description,
        })
        for column in table.columns:
            dictionary_rows.append({
                "table_key": table.key,
                "table_title": table.title,
                "column_key": column.key,
                "column_heading": column.heading,
                "unit": column.unit or "dimensionless / not applicable",
                "meaning": column.description or "Meaning follows the analysis result field name.",
            })
        if progress:
            progress(int(index * 75 / total), f"Writing TXT table {index}/{total}: {table.title}")

    readme_text = _clean_readme_text(package) if compact else _readme_text(package)
    readme = atomic_write_text(root / "README.txt", readme_text, backup=False)
    files.append(readme)
    if not compact:
        inventory = write_table_txt(
            root / "table_inventory.txt", inventory_rows,
            columns=["table_key", "table_title", "category", "dataset", "file", "rows", "columns", "description"],
            title="Raptor export table inventory", backup=False,
        )
        files.append(Path(inventory["txt_path"]))
        dictionary = write_table_txt(
            root / "data_dictionary.txt", dictionary_rows,
            columns=["table_key", "table_title", "column_key", "column_heading", "unit", "meaning"],
            title="Raptor export data dictionary", backup=False,
        )
        files.append(Path(dictionary["txt_path"]))
        metadata = write_table_txt(
            root / "package_metadata.txt",
            [{"parameter": key, "value": value} for key, value in package.metadata.items()],
            columns=["parameter", "value"], title="Raptor export package metadata", backup=False,
        )
        files.append(Path(metadata["txt_path"]))
    manifest_title = "Clean scientific data manifest" if compact else "Raptor export manifest"
    manifest = write_manifest_txt(root / "manifest.txt", files, root=root, title=manifest_title, backup=False)
    files.append(Path(manifest["txt_path"]))
    if progress:
        progress(100, "Clean TXT package complete")
    return root


def zip_directory(directory: str | Path, archive_path: str | Path) -> Path:
    root = Path(directory)
    target = Path(archive_path).with_suffix(".zip")
    target.parent.mkdir(parents=True, exist_ok=True)
    _backup_existing(target)
    descriptor, temp_name = tempfile.mkstemp(prefix=f".{target.name}.", suffix=".tmp.zip", dir=str(target.parent))
    os.close(descriptor)
    temporary = Path(temp_name)
    try:
        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for path in sorted(root.rglob("*")):
                if path.is_file():
                    archive.write(path, arcname=str(path.relative_to(root.parent)).replace("\\", "/"))
        os.replace(temporary, target)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    return target


def export_package_bundle(
    package: ExportPackage,
    output_directory: str | Path,
    *,
    base_name: str | None = None,
    excel: bool = True,
    text: bool = True,
    zip_text: bool = True,
    progress: Callable[[int, str], None] | None = None,
    cancelled: Callable[[], bool] | None = None,
    compact_text: bool = False,
) -> dict[str, Any]:
    package.validate()
    out = Path(output_directory)
    out.mkdir(parents=True, exist_ok=True)
    stem = safe_filename(base_name or package.project_name or "Afruz_PXRD_Export")
    result: dict[str, Any] = {"project_name": package.project_name, "table_count": len(package.tables)}
    stages = int(bool(excel)) + int(bool(text)) + int(bool(text and zip_text))
    completed = 0

    def stage_progress(local: int, message: str) -> None:
        if progress:
            start = int(completed * 100 / max(stages, 1))
            span = int(100 / max(stages, 1))
            progress(min(99, start + int(local * span / 100)), message)

    if excel:
        result["excel_path"] = str(write_excel_workbook(out / f"{stem}.xlsx", package, progress=stage_progress, cancelled=cancelled))
        completed += 1
    text_dir = None
    if text:
        text_dir = out / f"{stem}_TXT"
        if text_dir.exists():
            _backup_existing_directory(text_dir)
        result["text_directory"] = str(write_text_package(
            text_dir,
            package,
            progress=stage_progress,
            cancelled=cancelled,
            compact=compact_text,
        ))
        completed += 1
    if text_dir is not None and zip_text:
        if cancelled and cancelled():
            raise RuntimeError("Export cancelled by user.")
        if progress:
            progress(int(completed * 100 / max(stages, 1)), "Creating portable ZIP package")
        result["zip_path"] = str(zip_directory(text_dir, out / f"{stem}_TXT.zip"))
        completed += 1
    checksums = []
    for key in ("excel_path", "zip_path"):
        if result.get(key):
            p = Path(result[key])
            checksums.append({"file": p.name, "size_bytes": p.stat().st_size, "sha256": sha256_file(p)})
    if checksums:
        checksum_path = out / f"{stem}_checksums.txt"
        write_table_txt(
            checksum_path, checksums, columns=["file", "size_bytes", "sha256"],
            title="Export package checksums", backup=True,
        )
        result["checksums_path"] = str(checksum_path)
    if progress:
        progress(100, "Clean data export complete" if compact_text else "Raptor export complete")
    return result
