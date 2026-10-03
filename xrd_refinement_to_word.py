# -*- coding: utf-8 -*-
"""
Create a formatted Word document explaining the mathematics and logic of
Pawley, Le Bail, and Rietveld XRD refinement.

Install dependency:
    pip install python-docx

Run:
    python xrd_refinement_to_word.py
"""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


OUTPUT_FILE = "XRD_Refinement_Mathematics.docx"


def set_cell_shading(cell, fill="D9EAF7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_equation(doc, equation):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(equation)
    r.font.name = "Cambria Math"
    r.font.size = Pt(11)
    r.italic = True
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


doc = Document()

# ---------------- Page setup ----------------
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

# ---------------- Default fonts ----------------
styles = doc.styles
styles["Normal"].font.name = "Times New Roman"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
styles["Normal"].font.size = Pt(11)

for style_name in ["Title", "Heading 1", "Heading 2"]:
    styles[style_name].font.name = "Arial"
    styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

# ---------------- Title ----------------
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Mathematics and Logic of Powder XRD Refinement")
r.bold = True
r.font.name = "Arial"
r.font.size = Pt(18)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Pawley, Le Bail, and Rietveld Methods")
r.italic = True
r.font.size = Pt(12)

# ============================================================
doc.add_heading("1. General Principle of Powder XRD Refinement", level=1)

doc.add_paragraph(
    "All three methods attempt to make a calculated powder diffraction pattern "
    "agree with the experimentally measured XRD pattern. The main difference is "
    "how the integrated intensities of the Bragg reflections are determined."
)

add_equation(doc, "Measured pattern ≈ Calculated pattern")

doc.add_paragraph(
    "At each measured point i, usually corresponding to a particular 2θ value, "
    "the experiment provides an observed intensity yᵢ(obs), while the refinement "
    "model calculates yᵢ(calc). A simplified calculated powder pattern can be written as:"
)

add_equation(
    doc,
    "yᵢ(calc) = yᵢ(background) + Σₖ Iₖ Φ(2θᵢ − 2θₖ)"
)

add_bullets(doc, [
    "k: a crystallographic reflection such as (003), (006), or (012).",
    "2θₖ: calculated position of reflection k.",
    "Iₖ: integrated intensity of reflection k.",
    "Φ: peak-profile function.",
    "yᵢ(background): background contribution."
])

doc.add_paragraph(
    "The refinement normally minimizes a weighted least-squares objective function:"
)

add_equation(doc, "S = Σᵢ wᵢ [yᵢ(obs) − yᵢ(calc)]²")
add_equation(doc, "wᵢ ≈ 1 / σᵢ²")

doc.add_paragraph(
    "The model parameters are collected into a parameter vector p, and the optimizer "
    "searches for the parameter set that minimizes S."
)

add_equation(doc, "p = [p₁, p₂, …, pₙ]")
add_equation(doc, "p(best) = arg minₚ S")

# ============================================================
doc.add_heading("2. Origin of Peak Positions", level=1)

doc.add_paragraph(
    "The unit-cell geometry determines the allowed d-spacings. For a cubic system:"
)
add_equation(doc, "1/d²(hkl) = (h² + k² + l²) / a²")

doc.add_paragraph(
    "Other crystal systems use different metric equations, but the same principle applies. "
    "The d-spacing is related to diffraction angle through Bragg's law:"
)
add_equation(doc, "2 d(hkl) sin θ(hkl) = nλ")

doc.add_paragraph(
    "Therefore changes in a, b, c, α, β, and γ move the calculated peak positions. "
    "This is why Pawley and Le Bail refinements can determine accurate lattice parameters "
    "without refining atomic coordinates."
)

# ============================================================
doc.add_heading("3. Mathematical Description of Peak Shape", level=1)

doc.add_paragraph(
    "A real diffraction reflection has finite width and shape. Common profile functions "
    "include Gaussian, Lorentzian, pseudo-Voigt, Thompson-Cox-Hastings pseudo-Voigt, "
    "and fundamental-parameters models."
)

doc.add_paragraph("A simple pseudo-Voigt profile can be represented by:")
add_equation(doc, "Φ(x) = η L(x) + (1 − η) G(x)")

add_bullets(doc, [
    "L(x): Lorentzian component.",
    "G(x): Gaussian component.",
    "η: Lorentzian/Gaussian mixing parameter."
])

doc.add_paragraph(
    "A commonly used angular dependence of peak width is represented by the Caglioti relationship:"
)
add_equation(doc, "H² = U tan²θ + V tanθ + W")

doc.add_paragraph(
    "Here H is related to the peak full width at half maximum (FWHM), while U, V, and W "
    "describe profile broadening."
)

# ============================================================
doc.add_heading("4. Pawley Refinement", level=1)

doc.add_paragraph(
    "Pawley refinement is a whole-pattern method that refines the unit cell and profile "
    "without calculating Bragg intensities from an atomic structural model."
)

doc.add_paragraph(
    "Conceptually, the calculated pattern is:"
)
add_equation(doc, "yᵢ(calc) = bᵢ + Σ(hkl) I(hkl) Φ(hkl,i)")

doc.add_paragraph(
    "The integrated intensities I(hkl) are treated as adjustable quantities. "
    "Typical Pawley refinement parameters include:"
)

add_bullets(doc, [
    "Lattice parameters.",
    "Zero shift or specimen displacement.",
    "Background.",
    "Peak-profile parameters.",
    "Individual reflection intensities."
])

doc.add_paragraph(
    "Atomic coordinates are not required. Consequently, Pawley refinement is useful "
    "for checking whether a proposed unit cell and space group can reproduce the observed "
    "reflection positions and overall profile."
)

# ============================================================
doc.add_heading("5. Le Bail Refinement", level=1)

doc.add_paragraph(
    "Le Bail refinement has a similar purpose to Pawley refinement: it fits the whole "
    "pattern without requiring a complete atomic structural model. The main difference "
    "lies in the treatment of reflection intensities."
)

doc.add_paragraph(
    "Rather than treating every integrated intensity as an independent least-squares "
    "parameter, Le Bail refinement iteratively extracts and redistributes the observed "
    "intensity among overlapping reflections."
)

add_numbered(doc, [
    "Begin with estimated reflection intensities.",
    "Calculate the powder diffraction profile.",
    "Compare the calculated profile with the observed profile.",
    "Redistribute observed intensity among overlapping reflections.",
    "Update the extracted reflection intensities.",
    "Refine lattice and profile parameters.",
    "Repeat until convergence."
])

doc.add_paragraph("A conceptual update expression is:")
add_equation(
    doc,
    "I(hkl)⁽ⁿ⁺¹⁾ = I(hkl)⁽ⁿ⁾ × correction derived from [y(obs) / y(calc)]"
)

# ============================================================
doc.add_heading("6. Why Pawley and Le Bail Refinements Are Useful", level=1)

doc.add_paragraph(
    "These methods are particularly useful when the probable phase and approximate unit "
    "cell are known, but the detailed atomic arrangement is uncertain."
)

doc.add_paragraph("A logical workflow is:")
for line in [
    "Experimental XRD",
    "↓",
    "Background and profile model",
    "↓",
    "Pawley or Le Bail refinement",
    "↓",
    "Verify lattice parameters and peak indexing",
    "↓",
    "Check unexplained reflections",
    "↓",
    "Proceed to Rietveld refinement"
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(line)

# ============================================================
doc.add_heading("7. Rietveld Refinement", level=1)

doc.add_paragraph(
    "Rietveld refinement differs fundamentally from Pawley and Le Bail methods because "
    "the Bragg intensities are calculated from an explicit crystallographic structure."
)

doc.add_paragraph("A simplified Rietveld intensity equation is:")
add_equation(
    doc,
    "yᵢ(calc) = yᵢ(background) + s Σ(hkl) L(hkl) |F(hkl)|² Φ(hkl,i) P(hkl) A(hkl)"
)

add_bullets(doc, [
    "s: phase scale factor.",
    "L(hkl): Lorentz-polarization/geometrical contribution.",
    "F(hkl): crystallographic structure factor.",
    "Φ(hkl,i): peak-profile function.",
    "P(hkl): preferred-orientation correction.",
    "A(hkl): absorption or related correction."
])

doc.add_paragraph(
    "The key term is |F(hkl)|² because it links the diffraction intensity directly "
    "to the atomic structure."
)

# ============================================================
doc.add_heading("8. The Structure Factor", level=1)

doc.add_paragraph(
    "For a reflection hkl, the structure factor can be represented conceptually as:"
)

add_equation(
    doc,
    "F(hkl) = Σⱼ fⱼ oⱼ exp[−Bⱼ(sinθ/λ)²] exp[2πi(hxⱼ + kyⱼ + lzⱼ)]"
)

add_bullets(doc, [
    "fⱼ: atomic X-ray scattering factor.",
    "oⱼ: site occupancy.",
    "Bⱼ: atomic displacement/thermal parameter.",
    "xⱼ, yⱼ, zⱼ: fractional atomic coordinates."
])

add_equation(doc, "I(hkl) ∝ |F(hkl)|²")

doc.add_paragraph(
    "Consequently, changing atomic positions, occupancies, or displacement parameters "
    "can alter calculated relative peak intensities even when the unit-cell parameters "
    "remain unchanged."
)

# ============================================================
doc.add_heading("9. Physical Example", level=1)

doc.add_paragraph(
    "Consider two possible NiFe-LDH structural models with identical lattice parameters. "
    "Their peak positions may be almost identical because the unit cells are the same. "
    "A Pawley refinement can therefore fit their peak positions equally well."
)

doc.add_paragraph(
    "However, if the atomic arrangements differ, their structure factors F(003), F(006), "
    "F(012), and other reflections will differ. Therefore, the calculated relative "
    "intensities will also differ. Rietveld refinement can potentially discriminate "
    "between such models because it calculates intensities from the atomic structure."
)

# ============================================================
doc.add_heading("10. Pawley, Le Bail, and Rietveld Compared", level=1)

table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"

headers = ["Property", "Pawley", "Le Bail", "Rietveld"]
for i, text in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = text
    set_cell_shading(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in cell.paragraphs[0].runs:
        run.bold = True

rows = [
    ("Unit cell required", "Yes", "Yes", "Yes"),
    ("Space group required", "Usually", "Usually", "Yes"),
    ("Atomic coordinates required", "No", "No", "Yes"),
    ("Peak positions from crystallography", "Yes", "Yes", "Yes"),
    ("Reflection intensities", "Refined", "Iteratively extracted", "Calculated from structure"),
    ("Lattice refinement", "Yes", "Yes", "Yes"),
    ("Profile refinement", "Yes", "Yes", "Yes"),
    ("Atomic-position refinement", "No", "No", "Yes"),
    ("Occupancy refinement", "No", "No", "Possible"),
    ("Structural interpretation", "Limited", "Limited", "Yes"),
    ("Quantitative phase analysis", "Limited/indirect", "Limited/indirect", "Commonly used"),
]

for row_data in rows:
    cells = table.add_row().cells
    for i, value in enumerate(row_data):
        cells[i].text = value
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

# ============================================================
doc.add_heading("11. Why Rietveld Refinement Can Be Dangerous", level=1)

doc.add_paragraph(
    "Rietveld refinement frequently contains strongly correlated variables. "
    "Examples include:"
)

add_equation(doc, "Scale factor ↔ Occupancy")
add_equation(doc, "Occupancy ↔ Atomic displacement parameter")
add_equation(doc, "Crystallite size ↔ Microstrain")
add_equation(doc, "Zero shift ↔ Lattice parameters")
add_equation(doc, "Preferred orientation ↔ Structural intensity model")

doc.add_paragraph(
    "Therefore, decreasing the numerical residual does not automatically prove that "
    "the refined structural model is physically correct. A defensible refinement should "
    "combine good profile agreement with crystallographically reasonable parameters."
)

# ============================================================
doc.add_heading("12. Weighted Profile R-Factor (Rwp)", level=1)

doc.add_paragraph("A commonly reported profile residual is:")
add_equation(
    doc,
    "Rwp = √{ Σᵢ wᵢ[yᵢ(obs) − yᵢ(calc)]² / Σᵢ wᵢ[yᵢ(obs)]² } × 100"
)

doc.add_paragraph(
    "A lower Rwp generally indicates better profile agreement, but a lower Rwp does not "
    "necessarily imply a more physically correct crystal structure."
)

# ============================================================
doc.add_heading("13. Reduced Chi-Square", level=1)

doc.add_paragraph("A conceptual reduced chi-square expression is:")
add_equation(
    doc,
    "χ²red = [1 / (N − P)] Σᵢ {[yᵢ(obs) − yᵢ(calc)]² / σᵢ²}"
)

add_bullets(doc, [
    "N: number of observations.",
    "P: number of refined parameters.",
    "σᵢ: estimated uncertainty of observation i."
])

doc.add_paragraph(
    "A value near one can be meaningful when the uncertainty model is statistically "
    "appropriate. In practical powder diffraction, however, background errors, correlated "
    "data, imperfect counting statistics, instrument effects, and structural model errors "
    "can complicate interpretation."
)

# ============================================================
doc.add_heading("14. What the Numerical Optimizer Does", level=1)

doc.add_paragraph(
    "Suppose the current parameter vector contains lattice parameters, profile terms, "
    "structural coordinates, scale factors, and other variables:"
)
add_equation(doc, "p = [a, c, z(O), B(Ni), U, V, W, scale, …]")

doc.add_paragraph("At iteration n:")
add_equation(doc, "pₙ → yᵢ(calc)(pₙ)")
add_equation(doc, "rᵢ = yᵢ(obs) − yᵢ(calc)")
add_equation(doc, "Jᵢⱼ = ∂yᵢ(calc) / ∂pⱼ")
add_equation(doc, "pₙ₊₁ = pₙ + Δp")

doc.add_paragraph(
    "The derivatives form a Jacobian or design matrix. A least-squares algorithm estimates "
    "the parameter correction Δp, updates the parameters, recalculates the complete pattern, "
    "and repeats the cycle until a convergence criterion is reached."
)

# ============================================================
doc.add_heading("15. Recommended Logical Refinement Workflow", level=1)

workflow = [
    "Raw XRD data",
    "Instrument model",
    "Background",
    "Peak/profile assessment",
    "Phase identification",
    "Space group and approximate unit cell",
    "Pawley or Le Bail refinement",
    "Check whether the cell explains the peak positions",
    "Introduce CIF structural model",
    "Rietveld scale factor",
    "Lattice parameters",
    "Zero shift or specimen displacement",
    "Profile parameters",
    "Crystallite size and microstrain when justified",
    "Preferred orientation when justified",
    "Atomic coordinates",
    "Occupancy and displacement parameters only when supported",
    "Final physical and statistical validation",
]

for item in workflow:
    p = doc.add_paragraph(style="List Number")
    p.add_run(item)

# ============================================================
doc.add_heading("16. Short Summary", level=1)

p = doc.add_paragraph()
r = p.add_run("Pawley: ")
r.bold = True
p.add_run(
    "Fits the whole diffraction pattern without an atomic model; individual reflection "
    "intensities are refined."
)

p = doc.add_paragraph()
r = p.add_run("Le Bail: ")
r.bold = True
p.add_run(
    "Fits the whole diffraction pattern without an atomic model; reflection intensities "
    "are iteratively extracted from the measured profile."
)

p = doc.add_paragraph()
r = p.add_run("Rietveld: ")
r.bold = True
p.add_run(
    "Fits the diffraction pattern using an explicit crystallographic model, with reflection "
    "intensities calculated from the structure factor."
)

add_equation(
    doc,
    "Pawley / Le Bail: I(hkl) is obtained mainly from the observed diffraction pattern"
)
add_equation(
    doc,
    "Rietveld: I(hkl) ∝ |F(hkl)|² is calculated from the atomic structure"
)

doc.add_paragraph(
    "This distinction is the central conceptual difference between the three whole-pattern "
    "powder diffraction refinement approaches."
)

# ---------------- Save ----------------
doc.save(OUTPUT_FILE)
print(f"Created: {OUTPUT_FILE}")
